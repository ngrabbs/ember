"""Read-only extraction and independent numeric translation of TI calculator.
No workbook mutation/export or native Excel recalculation is performed.
"""
from pathlib import Path
import argparse,hashlib,json,math,warnings
import openpyxl
HERE=Path(__file__).resolve().parent

def startup(wb, vin=5, vout=5, c0=22, c5=10.5, load_type='Resistive Load', load=18, threshold=.5):
 prev=None; rows=[]
 for r in range(4,644):
  base=wb['Sheet2'][f'B{r}'].value
  if prev is None:
   control=6.5;trigger=0;duty=control;elapsed=25;charge=0;cap=c0;iload=0;vq=0
  else:
   control=prev['control'] if prev['elapsed']%4>0 else ((100 if prev['control']==100 else (prev['control']+6 if prev['control']==6.5 else prev['control']+12.5)) if prev['trigger']>0 else base)
   trigger=0 if prev['qvoltage']<2.7 else 1
   duty=(base if prev['qvoltage']<2.7 else control)*prev['running']
   elapsed=prev['elapsed']+25;charge=prev['charge'];vq=prev['qvoltage']
   cap=c0-vq*(c0-c5)/5
   iload=(vq*1000/load if load_type=='Resistive Load' else load)*(1 if vq>threshold else 0)
  supplied=(-.0457*vq+.64+(vin-5)*.148)*duty*10
  charge+=(supplied-iload)*25/1000
  qvoltage=charge/cap*(1 if prev is None else prev['running'])
  output=min(qvoltage,vout)
  running=(0 if output==vout else 1) if prev is None else (0 if prev['running']==0 or output==vout else 1)
  item={'control':control,'trigger':trigger,'duty':duty,'elapsed':elapsed,'charge':charge,'qvoltage':qvoltage,'output':output,'running':running,'supply_mA':supplied,'load_mA':iload}
  rows.append(item);prev=item
 return rows

def burst(wb,vin,cap,esr,load):
 # Tonmax!H2 approximate-match VLOOKUP, preserving native listed Vin bins.
 table=[(wb['Tonmax'][f'B{r}'].value,wb['Tonmax'][f'D{r}'].value) for r in range(27,35)]
 # D column is formula; numeric values taken from the source's cached companion.
 iburst=max((x for x in table if x[0]<=vin),key=lambda x:x[0])[1]
 vh=.010/.2415;charge=(iburst-.001*load)*286e-9+.001*load*222e-9
 ripple=vh+charge/(cap*1e-6)+iburst*esr*.001
 ton=1e6*cap*1e-6*ripple/(iburst-(load+2)*.001)
 return {'burst_A':iburst,'ripple_estimate_mV':ripple*1000,'ton_us':ton}

def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('--source',type=Path,required=True);a=p.parse_args()
 digest=hashlib.sha256(a.source.read_bytes()).hexdigest()
 assert digest=='6001da9f256856a29aa73adbbccd0972da93cc199bcabfb3478e128e3e6531e3','Unexpected TI calculator revision'
 with warnings.catch_warnings():
  warnings.simplefilter('ignore',UserWarning)
  f=openpyxl.load_workbook(a.source,data_only=False,read_only=False)
  cached=openpyxl.load_workbook(a.source,data_only=True,read_only=False)
 assert f.sheetnames==['Inputs','Sheet1','Sheet2','Tonmax']
 assert f['Inputs']['C21'].value=='=IF(Sheet2!Y2<16000, Sheet2!Y2, "Failed Startup")'
 default=startup(f)
 # Independent recomputation of all640 Q/S/I cells against shipped default cache.
 maxerr={key:max(abs(x[key]-cached['Sheet2'][f'{col}{r}'].value) for r,x in enumerate(default,4)) for key,col in [('qvoltage','Q'),('output','S'),('supply_mA','I')]}
 assert max(maxerr.values())<1e-9,maxerr
 rise=next(x['elapsed'] for x in default if x['output']==5)
 assert rise==cached['Inputs']['C21'].value==625
 default_burst=burst(cached,5,10.5,3,200)
 assert math.isclose(default_burst['ton_us'],cached['Inputs']['C23'].value,rel_tol=1e-12)
 cases=[]
 for vin in [3.0,3.1,3.3]:
  for c0,c5 in [(22,10.5),(33,16),(47,33)]:
   for load in [13,50,100]:
    rs=startup(f,vin=vin,c0=c0,c5=c5,load_type='Constant Current Load',load=load,threshold=0)
    reached=next((x['elapsed'] for x in rs if x['output']==5),None)
    cases.append({'vin_V':vin,'C0_uF':c0,'C5_uF':c5,'constant_load_mA':load,'load_threshold_V':0,'rise_us':reached,'before_16ms':reached is not None and reached<16000,'steady_state_13mA':burst(cached,vin,c5,3,13)})
 result={'source_url':'https://dr-download.ti.com/design-tools-simulation/calculation-tool/MD-MKWZTR3T9i/V1/UCC3302X-CALC-V1.xlsx','source_sha256':digest,'source_bytes':a.source.stat().st_size,'method':'independent numeric translation of selected TI formulas; source read-only; no Excel-engine recalc','cached_default_agreement_max_abs_errors':maxerr,'default_rise_us':rise,'default_burst':default_burst,'cases':cases,'cautions':['Calculator aid/typical empirical model, not a TI guaranteed limit','No prebias, repeated UVLO, coupled essential-source or LDO transient model','Capacitance pairs and13/50/100mA loads are sensitivity inputs, not qualified selected-part limits','Input Vin bins preserve approximate VLOOKUP; no interpolation claimed','Sheet2 U4 contains an unused self-reference; excluded from the translated rise calculation; W4=1 for these inputs becauseV4=1','Tonmax cold-current rows are model data, not datasheet warranty; do not promote them to guaranteed cold rating']}
 (HERE/'calculated.json').write_text(json.dumps(result,indent=2)+'\n')
 print({'source_sha256':digest,'cached_default_errors':maxerr,'cases':len(cases),'startup_failures':sum(not x['before_16ms'] for x in cases),'rise_us_range':[min(x['rise_us'] for x in cases if x['rise_us']),max(x['rise_us'] for x in cases if x['rise_us'])]})
if __name__=='__main__':main()
